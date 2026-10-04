#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a local Compendex snapshot index from Elsevier's official XLSX.

The generated file is intentionally NOT committed/distributed by RankLens.
Users obtain the official workbook themselves and run this script locally:

  python build/gen_ei_data.py path/to/COMPENDEX_Source-list-082026.xlsx

Official source page:
  https://www.elsevier.com/products/engineering-village/databases/compendex

The script records only normalized lookup keys and compact metadata needed to
show "listed in the official source snapshot". It does not assert that any
individual paper was indexed by Compendex.
"""

import argparse
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "data" / "eiIndex.js"
OFFICIAL_PAGE = "https://www.elsevier.com/products/engineering-village/databases/compendex"


def norm_title(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.upper().replace("&", " AND ")
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", value).split())


def norm_id(value):
    value = re.sub(r"[^0-9Xx]", "", str(value or ""))
    return value.upper() if value and value != "" else ""


def sheet_date(ws):
    text = str(ws.cell(1, 1).value or "")
    match = re.search(r"(?:UPDATED\s+)+(.*)$", text, re.I)
    return match.group(1).strip() if match else ""


def merge_record(mapping, key, record):
    """Merge duplicate source rows without hiding active/discontinued conflicts."""
    if not key:
        return
    existing = mapping.get(key)
    if existing is None:
        mapping[key] = record
        return
    old_status = existing[3] if len(existing) > 3 else ""
    new_status = record[3] if len(record) > 3 else ""
    if old_status != new_status:
        existing[3] = "conflict"
        if len(record) > 4 and len(existing) < 5:
            existing.append(record[4])


def add_serial(index, title, source_type, issn, eissn, sheet, status, coverage=None):
    title_key = norm_title(title)
    if not title_key:
        return
    record = [title, source_type or "", sheet, status]
    if coverage:
        record.append(coverage)
    merge_record(index["serialByName"], title_key, record)
    for raw in (issn, eissn):
        key = norm_id(raw)
        if key and key != "-":
            merge_record(index["serialByIssn"], key, record)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path, help="Official Compendex Source List XLSX")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    blob = args.xlsx.read_bytes()
    sha256 = hashlib.sha256(blob).hexdigest()
    wb = load_workbook(args.xlsx, read_only=True, data_only=True)

    index = {
        "serialByIssn": {},
        "serialByName": {},
        "proceedingByIsbn": {},
        "proceedingByName": {},
        "meta": {
            "officialPage": OFFICIAL_PAGE,
            "sourceFile": args.xlsx.name,
            "sha256": sha256,
            "builtAt": datetime.now(timezone.utc).isoformat(),
            "sheets": {},
        },
    }

    ws = wb["SERIALS"]
    index["meta"]["sheets"]["SERIALS"] = sheet_date(ws)
    for row in ws.iter_rows(min_row=3, values_only=True):
        add_serial(index, row[0], row[1], row[2], row[3], "SERIALS", "active")

    ws = wb["CHINESE JRS on SERIALS LIST"]
    index["meta"]["sheets"]["CHINESE JRS on SERIALS LIST"] = sheet_date(ws)
    for row in ws.iter_rows(min_row=3, values_only=True):
        title = row[4] or row[3] or row[2]
        add_serial(index, title, "Journal", row[0], row[1], "CHINESE JRS", str(row[6] or "active"))

    ws = wb["DISCONTINUED"]
    index["meta"]["sheets"]["DISCONTINUED"] = sheet_date(ws)
    for row in ws.iter_rows(min_row=4, values_only=True):
        coverage = {"year": row[4], "volume": row[5], "issue": row[6], "pages": row[7]}
        add_serial(index, row[0], "Serial", row[1], row[2], "DISCONTINUED", "discontinued", coverage)

    ws = wb["NON-SERIALS"]
    index["meta"]["sheets"]["NON-SERIALS"] = sheet_date(ws)
    for row in ws.iter_rows(min_row=3, values_only=True):
        title = str(row[0] or "").strip()
        if not title:
            continue
        record = [title, row[1] or "Proceeding", "NON-SERIALS", "listed"]
        title_key = norm_title(title)
        index["proceedingByName"].setdefault(title_key, record)
        isbn = norm_id(row[2])
        if isbn and isbn != "-":
            index["proceedingByIsbn"].setdefault(isbn, record)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(index, ensure_ascii=False, separators=(",", ":"))
    args.output.write_text(
        "/** Generated locally from Elsevier's official Compendex Source List.\n"
        " * Do not redistribute without permission from the data owner.\n"
        " */\nccf.eiIndex = " + payload + ";\n",
        encoding="utf-8",
    )
    print(f"wrote {args.output}")
    print(f"serial ISSNs: {len(index['serialByIssn'])}")
    print(f"serial names: {len(index['serialByName'])}")
    print(f"proceeding ISBNs: {len(index['proceedingByIsbn'])}")
    print(f"proceeding names: {len(index['proceedingByName'])}")
    print(f"sha256: {sha256}")


if __name__ == "__main__":
    main()
